"""Build the Southeast Asian Journal of Economics (SAJE) revision package for manuscript 292892.

Sources: manuscript/manuscript_saje.md (text, {{key}} numbers, exhibits), project_R/outputs/numbers.csv and
numbers_saje.csv (numbers), project_R/outputs/*_formatted.csv (tables), outputs/figures (figures).
Marking: a sentence is printed in blue in the marked files when it does not appear verbatim in the round-0 submission
(ORIG_DOCX), and text wrapped in {+ ... +} is always blue. The clean files print everything in black.
SAJE guidelines: blinded article, separate title page, abstract of 150 words or less, tables and figures within the
text at the appropriate points. Section numbers follow the journal's sample paper ("1. Introduction", "2.1 Data").
Usage: python3 build_saje.py OUTDIR
"""
import csv, os, re, subprocess, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_manuscript import (blocks, base_doc, set_spacing, latex_to_omml, save_doc, border, OUTR, ROOT)
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

MS = os.path.join(ROOT, 'manuscript', 'manuscript_saje.md')
LETTER = os.path.join(ROOT, 'manuscript', 'response_letter_saje.md')
SRC0 = '/home/user/trungcandygit/b-i-esg2/submission/AAMJAF_submission'
ORIG_DOCX = [os.path.join(SRC0, '02_Manuscript_Anonymized.docx'), os.path.join(SRC0, '05_Supplementary_material.docx')]
BLUE = RGBColor(0, 0, 0xFF)
TOKEN = re.compile(r'(\*[^*]+\*|\^[^^]+\^|~[^~]+~)')
STAR = '∗'
REF = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'saje_reference.docx')


def load_numbers():
    nums = {}
    for f in ('numbers.csv', 'numbers_saje.csv'):
        with open(os.path.join(OUTR, f), encoding='utf-8') as fh:
            for r in csv.DictReader(fh):
                assert r['key'] not in nums or nums[r['key']] == r['value'], 'duplicate key ' + r['key']
                nums[r['key']] = r['value']
    return nums


def fill(text, nums):
    def rep(m):
        if m.group(1) not in nums:
            sys.exit('Unknown placeholder {{%s}}' % m.group(1))
        return nums[m.group(1)]
    return re.sub(r'\{\{([A-Za-z0-9_]+)\}\}', rep, text)


def parse(path, nums):
    secs, cur = {}, None
    for line in open(path, encoding='utf-8').read().splitlines():
        if line.startswith('%'):
            continue
        if line.startswith('@@ '):
            cur = line[3:].strip(); secs[cur] = []; continue
        if cur is not None:
            secs[cur].append(fill(line, nums))
    return {k: '\n'.join(v).strip() for k, v in secs.items()}


def strip_marks(t):
    return t.replace('{+', '').replace('+}', '')


# ---- automatic change marking against the round-0 submission ----
def norm(s):
    s = strip_marks(s).replace('\\*', '*')
    s = re.sub(r'[*^~]', '', s)
    s = s.replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"').replace(' ', ' ')
    return re.sub(r'\s+', ' ', s).strip()


SPLIT = re.compile(r'(?<=[.?!:;])\s+(?=[A-Z(\[0-9H])')


def orig_sentences():
    out = set()
    for f in ORIG_DOCX:
        d = Document(f)
        texts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                texts += [c.text for c in row.cells]
        for tx in texts:
            for s in SPLIT.split(tx):
                n = norm(s)
                if n:
                    out.add(n)
    return out


ORIG = orig_sentences()


def automark(text):
    """Wrap every sentence that is not in the round-0 submission in {+ +}; keep explicit marks."""
    parts = SPLIT.split(text)
    out = []
    for s in parts:
        if '{+' in s or '+}' in s:
            out.append(s)   # explicitly marked already
        elif norm(s) and norm(s) not in ORIG:
            out.append('{+' + s + '+}')
        else:
            out.append(s)
    return ' '.join(out)


class Writer:
    """Adds runs with italics, super- and subscripts, and blue for {+ ... +} when marked=True."""

    def __init__(self, marked):
        self.marked = marked

    def runs(self, p, text, bold=False, size=None, italic=False, force_blue=False, auto=True):
        if auto and self.marked and not force_blue:
            text = automark(text)
        text = text.replace('\\*', STAR).replace('\n', ' ')
        blue = False
        for seg in re.split(r'(\{\+|\+\})', text):
            if seg == '{+':
                blue = True; continue
            if seg == '+}':
                blue = False; continue
            for part in TOKEN.split(seg):
                if not part:
                    continue
                r = p.add_run()
                if part.startswith('*') and part.endswith('*') and len(part) > 2:
                    r.text = part[1:-1]; r.italic = True
                elif part.startswith('^') and part.endswith('^') and len(part) > 2:
                    r.text = part[1:-1]; r.font.superscript = True
                elif part.startswith('~') and part.endswith('~') and len(part) > 2:
                    r.text = part[1:-1]; r.font.subscript = True
                else:
                    r.text = part
                r.text = r.text.replace(STAR, '*')
                r.bold = bold or None
                if italic:
                    r.italic = True
                if size:
                    r.font.size = Pt(size)
                if self.marked and (blue or force_blue):
                    r.font.color.rgb = BLUE
        return p


def sec_number(text):
    """'## 2.1. Title' in the source becomes '2.1 Title' (journal style); level-1 keeps '1. Title'."""
    return re.sub(r'^(\{\+)?(\d+\.\d+(?:\.\d+)?)\.\s', lambda m: (m.group(1) or '') + m.group(2) + ' ', text)


def heading(d, w, text, level):
    p = d.add_paragraph(style='Heading %d' % level)
    w.runs(p, sec_number(text), bold=True, size=13 if level == 1 else 12)
    for r in p.runs:   # real Word headings, typeset like the journal's sample (Times New Roman, black)
        r.font.name = 'Times New Roman'; r.italic = False
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
        if r.font.color.rgb != BLUE:
            r.font.color.rgb = RGBColor(0, 0, 0)
    set_spacing(p, 1.5, 6, 18 if level == 1 else 8)
    p.paragraph_format.keep_with_next = True
    return p


def body_par(d, w, text, indent=True):
    p = d.add_paragraph(); w.runs(p, text); set_spacing(p, 1.5, 6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    return p


def equation(d, w, text):
    blue = text.startswith('{+')
    text = strip_marks(text)
    eq, num = [s.strip() for s in text[len('$$latex'):].rsplit('|', 1)]
    p = d.add_paragraph(); set_spacing(p, 1.5, 6, 6)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t'); p._p.append(latex_to_omml(eq)); r = p.add_run('\t' + num)
    if blue and w.marked:
        r.font.color.rgb = BLUE
    return p


def set_cell(w, cell, text, bold=False, align=None, size=10, blue=False):
    cell.text = ''
    p = cell.paragraphs[0]; w.runs(p, text, bold=bold, size=size, force_blue=blue, auto=False)
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align


def table_exhibit(d, w, caption, csvfile, note, source, widths=None, size=9, blue=False):
    rows = list(csv.reader(open(os.path.join(OUTR, csvfile), encoding='utf-8')))
    lab, title = [s.strip() for s in caption.split('|')]
    p = d.add_paragraph(); w.runs(p, lab + '. ' + title, force_blue=blue, size=11)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 1.0, 6, 12); p.paragraph_format.keep_with_next = True
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 30), ('bottom', 30), ('left', 70), ('right', 70)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    t._tbl.tblPr.append(mar)
    for i, row in enumerate(rows):
        trPr = t.rows[i]._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit'); trPr.append(cs)
        if i == 0:   # header row repeats on a new page
            th = OxmlElement('w:tblHeader'); trPr.append(th)
        for j, v in enumerate(row):
            c = t.cell(i, j)
            set_cell(w, c, v, bold=(i == 0), align=None if j == 0 else WD_ALIGN_PARAGRAPH.CENTER, size=size, blue=blue)
            if i < len(rows) - 1:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True   # keep the table on one page where it fits
            if i == 0:
                border(c, ['top', 'bottom'])
            if i == len(rows) - 1:
                border(c, ['bottom'])
    if widths:
        for j, wd in enumerate(widths):
            for i in range(len(rows)):
                t.cell(i, j).width = Inches(wd)
        for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(wd * 1440)))
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    q = d.add_paragraph(); w.runs(q, 'Note: ' + note, size=9, force_blue=blue); set_spacing(q, 1.0, 2, 4)
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q = d.add_paragraph(); w.runs(q, 'Source: ' + source, size=9, force_blue=blue); set_spacing(q, 1.0, 12)


def figure_exhibit(d, w, caption, png, note, source, blue=False, width=6.5):
    lab, title = [s.strip() for s in caption.split('|')]
    p = d.add_paragraph(); w.runs(p, lab + '. ' + title, force_blue=blue, size=11)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 1.0, 6, 12); p.paragraph_format.keep_with_next = True
    d.add_picture(os.path.join(OUTR, 'figures', png), width=Inches(float(width)))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; d.paragraphs[-1].paragraph_format.keep_with_next = True
    q = d.add_paragraph(); w.runs(q, 'Note: ' + note, size=9, force_blue=blue); set_spacing(q, 1.0, 2, 4)
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q = d.add_paragraph(); w.runs(q, 'Source: ' + source, size=9, force_blue=blue); set_spacing(q, 1.0, 12)


def parse_exhibits(text):
    exs = []
    for b in blocks(text):
        kv = dict(line.split(':=', 1) for line in b.splitlines())
        kv = {k.strip(): v.strip() for k, v in kv.items()}
        kind = kv.pop('kind')
        if 'widths' in kv:
            kv['widths'] = [float(x) for x in kv['widths'].split(',')]
        if 'size' in kv:
            kv['size'] = int(kv['size'])
        kv['blue'] = kv.get('blue', '0') == '1'
        exs.append({'kind': kind, 'args': kv})
    return exs


def place(d, w, ex):
    (table_exhibit if ex['kind'] == 'table' else figure_exhibit)(d, w, **ex['args'])


def write_body(d, w, text, exhibits):
    left = list(exhibits); pending = []
    bl = blocks(text)
    for k, b in enumerate(bl):
        plain = strip_marks(b)
        if plain.startswith('## '):
            heading(d, w, b.replace('## ', '', 1), 2)
        elif plain.startswith('# '):
            heading(d, w, b.replace('# ', '', 1), 1)
        elif plain.startswith('$$latex'):
            equation(d, w, b)
        else:
            body_par(d, w, b)
            for ex in list(left):   # SAJE: each exhibit after the paragraph that first cites it
                lab = ex['args']['caption'].split('|')[0].strip()
                if re.search(r'\b' + re.escape(lab) + r'(?![\dA-Za-z])', plain):
                    pending.append(ex); left.remove(ex)
        nxt = strip_marks(bl[k + 1]) if k + 1 < len(bl) else ''
        if pending and not nxt.startswith('$$latex') and not plain.rstrip().endswith((' is', ':')):
            for ex in pending:   # never between a sentence and the equation that completes it
                place(d, w, ex)
            pending = []
    for ex in pending:
        place(d, w, ex)
    assert not left, 'exhibit never cited: ' + str([e['args']['caption'] for e in left])


def references(d, w, text):
    heading(d, w, 'References', 1)
    for b in blocks(text):
        p = d.add_paragraph(); w.runs(p, b); set_spacing(p, 1.5, 4)
        p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)


def words(t):
    t = strip_marks(t)
    t = re.sub(r'\$\$latex.*', '', t)
    return len(re.findall(r'\S+', re.sub(r'^#+ .*$', '', t, flags=re.M)))


def manuscript(out, secs, marked):
    w = Writer(marked); d = base_doc()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(1)
    p = d.add_paragraph(); w.runs(p, secs['title'], bold=True, size=14); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, 1.5, 18)
    p = d.add_paragraph(); w.runs(p, 'Abstract', bold=True, size=13); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, 1.5, 6)
    body_par(d, w, secs['abstract'], indent=True)
    p = d.add_paragraph(); w.runs(p, 'Keywords:', bold=True, auto=False); w.runs(p, ' ' + secs['keywords'])
    set_spacing(p, 1.5, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p = d.add_paragraph(); w.runs(p, '{+JEL Classification:+}', bold=True, auto=False); w.runs(p, ' {+' + secs['jel'] + '+}')
    set_spacing(p, 1.5, 6)
    d.add_page_break()
    write_body(d, w, secs['body'], parse_exhibits(secs['exhibits']))
    references(d, w, secs['references'])
    save_doc(d, out)


def supplement(out, secs, marked):
    w = Writer(marked); d = base_doc()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(1)
    p = d.add_paragraph(); w.runs(p, 'Supplementary material', bold=True, size=14); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(); w.runs(p, secs['title'], bold=True); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 1.5, 12)
    body_par(d, w, secs['supp_intro'], indent=False)
    write_body(d, w, secs['supp_body'], [])
    for ex in parse_exhibits(secs['exhibits_supp']):   # all exhibits after the appendix text, in numbered order
        place(d, w, ex)
    save_doc(d, out)


def pandoc(md_text, out):
    md = out + '.md'
    open(md, 'w', encoding='utf-8').write(md_text + '\n')
    subprocess.run(['pandoc', md, '-f', 'markdown', '-o', out, '--reference-doc', REF], check=True)
    os.remove(md)


def reference_doc():
    """Pandoc reference document with Times New Roman 12 pt and 1.5 line spacing."""
    data = subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'], check=True,
                          capture_output=True).stdout
    open(REF, 'wb').write(data)
    d = Document(REF)
    for st in d.styles:
        try:
            if st.font is not None:
                st.font.name = 'Times New Roman'
                st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
                if st.name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
                    st.font.size = Pt(12)
                if st.name.startswith('Heading') or st.name == 'Title':
                    st.font.color.rgb = RGBColor(0, 0, 0)
        except (AttributeError, ValueError):
            pass
    for name in ('Body Text', 'First Paragraph'):
        d.styles[name].paragraph_format.line_spacing = 1.5
    d.save(REF)


def letters(outdir, nums, secs):
    """Two response files (editor; Reviewer A), so that the referee's copy holds only the referee's part."""
    n2 = dict(nums)
    n2['s_abstract_text'] = strip_marks(secs['abstract']); n2['s_abstract_words'] = str(words(secs['abstract']))
    L = parse(LETTER, n2)
    pandoc(L['editor'], os.path.join(outdir, '06_Response_to_Editor.docx'))
    pandoc(L['reviewer'], os.path.join(outdir, '07_Response_to_Reviewer_A.docx'))
    return L


if __name__ == '__main__':
    outdir = sys.argv[1]
    os.makedirs(outdir, exist_ok=True)
    nums = load_numbers(); secs = parse(MS, nums)
    reference_doc()
    manuscript(os.path.join(outdir, '02_Manuscript_Revised_Marked_blue.docx'), secs, True)
    manuscript(os.path.join(outdir, '03_Manuscript_Revised_Clean.docx'), secs, False)
    supplement(os.path.join(outdir, '04_Supplementary_Material_Marked_blue.docx'), secs, True)
    supplement(os.path.join(outdir, '05_Supplementary_Material_Clean.docx'), secs, False)
    pandoc(secs['title_page'], os.path.join(outdir, '01_Title_Page.docx'))
    L = letters(outdir, nums, secs)
    ab = words(secs['abstract']); body = words(secs['body'])
    print('abstract words:', ab, '| body words (excl. headings, equations):', body,
          '| letter words: editor', len(L['editor'].split()), 'reviewer', len(L['reviewer'].split()))
    assert ab <= 150, 'abstract over 150 words'
    for s in ('body', 'abstract'):
        for k in ('{{', '[[', 'AUTHOR_INPUT'):
            assert k not in secs[s], 'unresolved %s in %s' % (k, s)
